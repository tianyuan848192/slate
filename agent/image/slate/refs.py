"""参考资料（docs/references-design.md）。

slate-refs build   开工前由 turn Lambda 在会话里执行（不经过模型）：把 input/refs/ 下的每份资料抽成纯文本 input/refs/<编号>.md，
                   生成目录 input/refs/index.json（编号、标题、类型、备注、来源、字数、章节标题和行号）。
                   资料清单 input/refs/manifest.json 由 api Lambda 建任务时写好；链接的正文由开工前的抓取放在 input/refs/web/<编号>.txt。
slate-refs show    打印目录（Agent 用）

剧本核对 check_script() 由 slate-publish script 调用：交给 Slate-llm（Sonnet），比对剧本和资料。
"""
import argparse, csv, io, json, os, re, sys

from .common import ctx, die, load_json, path

READ_ALL = 30000      # 一份资料不超过这么多字时，要求 Agent 读全文
TOTAL_LIMIT = 200000  # 所有资料合计的字数上限（超出的部分截断，并在目录里标出来）
CHECK_LIMIT = 120000  # 核对时交给模型的资料字数上限
TABLE_ROWS = 200


def _read_text(p):
    raw = open(p, "rb").read()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _table_md(rows):
    rows = [["" if v is None else str(v).replace("|", "\\|").replace("\n", " ") for v in r] for r in rows if any(v not in (None, "") for v in r)]
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [r + [""] * (w - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * w]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:TABLE_ROWS]]
    if len(rows) > TABLE_ROWS:
        out.append(f"（共 {len(rows) - 1} 行，这里只列前 {TABLE_ROWS - 1} 行；完整数据在原文件里）")
    return "\n".join(out)


def extract(p):
    """一个文件 → 纯文本（表格转成 Markdown 表格）"""
    ext = os.path.splitext(p)[1].lower()
    if ext in (".md", ".txt"):
        return _read_text(p)
    if ext == ".csv":
        return _table_md(list(csv.reader(io.StringIO(_read_text(p)))))
    if ext == ".pdf":
        from pypdf import PdfReader
        pages = PdfReader(p).pages
        return "\n\n".join(f"## 第 {i} 页\n\n{(pg.extract_text() or '').strip()}" for i, pg in enumerate(pages, 1))
    if ext == ".docx":
        import docx
        d = docx.Document(p)
        out = []
        for para in d.paragraphs:
            t = para.text.strip()
            if not t:
                continue
            lvl = re.match(r"Heading (\d)", para.style.name or "")
            out.append(("#" * (int(lvl.group(1)) + 1) + " " + t) if lvl else t)
        for i, tb in enumerate(d.tables, 1):
            out.append(f"\n## 表格 {i}\n\n" + _table_md([[c.text for c in r.cells] for r in tb.rows]))
        return "\n\n".join(out)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        return "\n\n".join(f"## 工作表 {ws.title}\n\n" + _table_md(list(ws.iter_rows(values_only=True))) for ws in wb.worksheets)
    raise ValueError(f"不支持的格式：{ext}")


def _headings(text):
    out = []
    for i, line in enumerate(text.splitlines(), 1):
        if re.match(r"#{1,4} \S", line):
            out.append([i, line.lstrip("#").strip()[:60]])
    return out[:80]


def build():
    man = load_json(path("input", "refs", "manifest.json"))
    if not man:
        return {"items": [], "totalChars": 0}
    items, total = [], 0
    for it in man.get("items", []):
        rid = it["id"]
        rec = {k: it.get(k) for k in ("id", "kind", "title", "note", "source")}
        try:
            if it["kind"] == "link":
                p = path("input", "refs", "web", f"{rid}.txt")
                if not os.path.isfile(p):
                    raise ValueError("网页打不开，没有抓到正文")
                text = _read_text(p)
            else:
                text = extract(path(it["file"]))
        except Exception as e:
            rec.update(error=str(e)[:200], chars=0)
            items.append(rec)
            continue
        text = re.sub(r"\n{3,}", "\n\n", text).strip()
        if total + len(text) > TOTAL_LIMIT:
            text = text[:max(0, TOTAL_LIMIT - total)]
            rec["truncated"] = True
        total += len(text)
        head = f"# {it.get('title')}\n\n> 类型：{ {'file': '文件', 'text': '粘贴的文字', 'link': '链接'}[it['kind']] }；来源：{it.get('source')}" + \
               (f"；用户备注：{it['note']}" if it.get("note") else "") + "\n\n"
        open(path("input", "refs", f"{rid}.md"), "w", encoding="utf-8").write(head + text)
        rec.update(path=f"input/refs/{rid}.md", chars=len(text), readAll=len(text) <= READ_ALL, headings=_headings(head + text))
        items.append(rec)
    index = {"strict": bool(man.get("strict")), "web": bool(man.get("web")), "totalChars": total, "readAllUnder": READ_ALL, "items": items}
    json.dump(index, open(path("input", "refs", "index.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return index


def index():
    return load_json(path("input", "refs", "index.json"))


CHECK_SYSTEM = ("你是讲解视频和广告视频的事实核对员。只根据给你的参考资料判断，不用自己的知识补充。"
                "逐句检查剧本里的事实、数字、专业术语：和资料矛盾的标 contradiction；资料里找不到依据的标 unsupported；"
                "术语的用法和资料里的解释不一致的标 term。修辞、比喻、观点、过渡句和常识性的话不用标。"
                "只输出一个 JSON 对象，不要输出别的：{\"issues\": [{\"kind\": \"contradiction|unsupported|term\", "
                "\"text\": \"剧本原句\", \"ref\": \"资料编号\", \"suggestion\": \"怎么改\"}]}。没有问题时 issues 为空数组。")


def check_script(script):
    """剧本核对：返回 {"ok", "strict", "issues", "blocking"}；没有资料或核对调不通时返回 None"""
    idx = index()
    if not idx or not [x for x in idx["items"] if x.get("path")]:
        return None
    refs, n = [], 0
    for x in idx["items"]:
        if not x.get("path"):
            continue
        text = open(path(x["path"]), encoding="utf-8").read()
        text = text[:max(0, CHECK_LIMIT - n)]
        n += len(text)
        refs.append(f"<资料 编号=\"{x['id']}\">\n{text}\n</资料>")
    paras = "\n".join(f"{i}. {p}" for i, p in enumerate(script.get("paragraphs") or [], 1))
    question = "参考资料：\n" + "\n\n".join(refs) + f"\n\n剧本（逐段）：\n{paras}\n\n请核对。"
    import boto3
    from .common import REGION
    fn = os.environ.get("SLATE_LLM_FN")
    if not fn:
        return None
    try:
        payload = json.dumps({"op": "text", "images": [], "question": question, "system": CHECK_SYSTEM, "maxTokens": 8000,
                              "billing": {"taskId": ctx().get("taskId"), "chapter": ctx().get("chapter"), "token": ctx().get("billToken")}})
        from botocore.config import Config
        lam = boto3.client("lambda", region_name=REGION, config=Config(read_timeout=200, retries={"max_attempts": 1}))  # llm Lambda 最长 180 秒
        r = json.loads(lam.invoke(FunctionName=fn, Payload=payload.encode())["Payload"].read())
        m = re.search(r"\{[\s\S]*\}", r.get("text") or "")
        issues = json.loads(m.group(0)).get("issues", []) if (r.get("ok") and m) else None
    except Exception as e:
        print(json.dumps({"script_check": "unavailable", "detail": str(e)[:200]}, ensure_ascii=False), file=sys.stderr)
        return None
    if issues is None:
        return None
    issues = [{k: str(i.get(k) or "")[:300] for k in ("kind", "text", "ref", "suggestion")} for i in issues
              if i.get("kind") in ("contradiction", "unsupported", "term")][:30]
    strict = idx.get("strict")
    blocking = [i for i in issues if strict or i["kind"] == "contradiction"]
    return {"ok": not blocking, "strict": strict, "issues": issues, "blocking": len(blocking)}


def check_notes(notes):
    """资料摘要必须交代目录里的每一份资料"""
    idx = index() or {"items": []}
    want = [x["id"] for x in idx["items"]]
    got = {x.get("id") for x in notes.get("items", [])}
    miss = [i for i in want if i not in got]
    return [f"资料 {', '.join(miss)} 没有交代：每份资料都要写一段（用了什么，或者为什么没用）"] if miss else []


def main():
    ap = argparse.ArgumentParser(prog="slate-refs", description="参考资料：build 抽取文本并生成目录（开工前由系统执行）；show 打印目录")
    ap.add_argument("cmd", choices=["build", "show"])
    a = ap.parse_args()
    if a.cmd == "build":
        idx = build()
        print(json.dumps({"items": [{k: x.get(k) for k in ("id", "chars", "error", "truncated")} for x in idx["items"]],
                          "totalChars": idx.get("totalChars", 0)}, ensure_ascii=False))
    else:
        idx = index()
        if not idx:
            die("这个任务没有参考资料")
        print(json.dumps(idx, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
